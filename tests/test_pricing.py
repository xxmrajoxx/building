import io

import pandas as pd
import pytest
from openpyxl import load_workbook

from estimator import export, project_io
from estimator.estimate import EstimateSettings, build_lines, possible_duplicates, summarise
from estimator.ifc_takeoff import open_ifc, takeoff_ifc
from estimator.models import items_to_frame, new_item
from estimator.rates import auto_match, load_rates, unit_rate


@pytest.fixture(scope="module")
def rates():
    return load_rates()


@pytest.fixture(scope="module")
def ifc_takeoff(sample_ifc):
    items = takeoff_ifc(open_ifc(sample_ifc)).items
    return items[items["include"]].drop(columns="include").reset_index(drop=True)


def test_auto_match_ifc_items(ifc_takeoff, rates):
    matched = auto_match(ifc_takeoff, rates).set_index("description")["rate_code"]
    expect = {
        "Brick veneer": "MAS-001", "Stud partition": "PLA-002", "(volume)": "CON-001", "Slab - Ground slab": "CON",
        "Roof": "ROO-001", "Internal hollow core": "DOR-001", "External solid core": "DOR-002",
        "Aluminium window": "WIN-002", "Reinforcement": "CON-007",
    }
    for fragment, code in expect.items():
        codes = matched[matched.index.str.contains(fragment, regex=False)]
        assert len(codes) and all(c.startswith(code) for c in codes), (fragment, list(codes))


def test_no_match_when_units_differ(rates):
    df = items_to_frame([new_item("Masonry", "Face brickwork", "m3", 10, "Manual")])
    assert auto_match(df, rates)["rate_code"][0] == ""


def test_unit_rate_includes_waste():
    assert unit_rate({"material": 100, "waste_pct": 10, "labour": 50, "plant": 5}) == pytest.approx(165)


def test_markup_cascade():
    lines = pd.DataFrame({"trade": ["Concrete"], "amount": [100_000.0], "priced": [True]})
    s = EstimateSettings(preliminaries_pct=10, contingency_pct=5, escalation_pct=0, overhead_pct=5,
                         profit_pct=5, gst_pct=10)
    out = summarise(lines, s)
    # 100k + 10% prelims = 110k; +5% contingency = 115.5k; +10% OH&P = 127,050; +GST
    assert out["total_ex_gst"] == pytest.approx(127_050)
    assert out["total_inc_gst"] == pytest.approx(139_755)


def test_unpriced_and_location_factor(rates):
    df = items_to_frame([
        new_item("Masonry", "Brick", "m2", 10, "Manual", rate_code="MAS-001"),
        new_item("Masonry", "Brick wrong unit", "m3", 10, "Manual", rate_code="MAS-001"),
        new_item("Other", "Mystery", "item", 1, "Manual"),
    ])
    base = unit_rate(rates.set_index("code").loc["MAS-001"])
    lines = build_lines(df, rates, location_factor=1.1)
    priced = lines[lines["priced"]]
    assert len(priced) == 1
    assert priced["rate"].iloc[0] == pytest.approx(round(base * 1.1, 2))
    assert set(lines[~lines["priced"]]["issue"]) == {"no rate assigned", "unit mismatch (rate is per m2)"}


def test_possible_duplicates():
    df = items_to_frame([
        new_item("Doors & Hardware", "Door A", "no", 3, "IFC", rate_code="DOR-001"),
        new_item("Doors & Hardware", "Door A (schedule)", "no", 3, "PDF schedule", rate_code="DOR-001"),
        new_item("Doors & Hardware", "Door B", "no", 1, "IFC", rate_code="DOR-002"),
    ])
    assert len(possible_duplicates(df)) == 2


def test_excel_export_has_formulas(ifc_takeoff, rates):
    takeoff = auto_match(ifc_takeoff, rates)
    s = EstimateSettings()
    lines = build_lines(takeoff, rates)
    wb = load_workbook(io.BytesIO(export.to_excel({"name": "Test"}, s, lines, takeoff)))
    assert wb.sheetnames == ["Summary", "BoQ", "Takeoff"]
    formulas = [c.value for row in wb["BoQ"].iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert any(f.startswith("=E") and "*F" in f for f in formulas)
    summary_vals = [c.value for row in wb["Summary"].iter_rows() for c in row]
    assert "Tender price (incl. GST)" in summary_vals


def test_project_roundtrip(ifc_takeoff):
    s = EstimateSettings(location="Perth", profit_pct=8)
    raw = project_io.dump_project({"name": "X"}, s, ifc_takeoff)
    project, s2, t2 = project_io.load_project(raw)
    assert project["name"] == "X" and s2 == s
    assert t2["quantity"].sum() == pytest.approx(ifc_takeoff["quantity"].sum())
