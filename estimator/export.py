"""Excel export of the estimate: Summary, Bill of Quantities and Takeoff sheets.

Amounts and totals are written as formulas so an estimator can adjust rates or
markups in Excel and the tender price updates.
"""
from __future__ import annotations

import io
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .estimate import EstimateSettings

MONEY = '"$"#,##0.00'
QTY = "#,##0.00"
PCT = '0.0"%"'
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
TRADE_FILL = PatternFill("solid", fgColor="DCE6F1")
WARN_FILL = PatternFill("solid", fgColor="FCE4D6")
BOLD = Font(bold=True)
WHITE_BOLD = Font(bold=True, color="FFFFFF")
TOP = Border(top=Side(style="thin"))


def _header(ws, row: int, labels: list[str]) -> None:
    for col, label in enumerate(labels, 1):
        c = ws.cell(row=row, column=col, value=label)
        c.font, c.fill = WHITE_BOLD, HEADER_FILL
        c.alignment = Alignment(vertical="center")


def _widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def to_excel(project: dict, settings: EstimateSettings, lines: pd.DataFrame, takeoff: pd.DataFrame) -> bytes:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    boq = wb.create_sheet("BoQ")
    raw = wb.create_sheet("Takeoff")

    # ---- BoQ
    boq["A1"] = f"Bill of Quantities - {project.get('name') or 'Untitled project'}"
    boq["A1"].font = Font(bold=True, size=14)
    boq["A2"] = f"{project.get('client', '')}  {project.get('address', '')}".strip()
    boq["A3"] = f"Rates include location factor {settings.location} x{settings.location_factor:g}. Unpriced items highlighted."
    cols = ["Item", "Trade", "Description", "Unit", "Qty", "Rate", "Amount", "Rate code", "Source", "Reference"]
    _header(boq, 5, cols)
    r = 6
    trades = list(dict.fromkeys(lines["trade"])) if not lines.empty else []
    for t_idx, trade in enumerate(trades, 1):
        boq.cell(row=r, column=1, value=f"{t_idx}").font = BOLD
        c = boq.cell(row=r, column=3, value=trade.upper())
        c.font = BOLD
        for col in range(1, len(cols) + 1):
            boq.cell(row=r, column=col).fill = TRADE_FILL
        r += 1
        first = r
        for i, (_, line) in enumerate(lines[lines["trade"] == trade].iterrows(), 1):
            boq.cell(row=r, column=1, value=f"{t_idx}.{i}")
            boq.cell(row=r, column=2, value=trade)
            boq.cell(row=r, column=3, value=line["description"]).alignment = Alignment(wrap_text=True)
            boq.cell(row=r, column=4, value=line["unit"])
            boq.cell(row=r, column=5, value=float(line["quantity"])).number_format = QTY
            boq.cell(row=r, column=6, value=float(line["rate"])).number_format = MONEY
            boq.cell(row=r, column=7, value=f"=E{r}*F{r}").number_format = MONEY
            boq.cell(row=r, column=8, value=line["rate_code"])
            boq.cell(row=r, column=9, value=line["source"])
            boq.cell(row=r, column=10, value=line["ref"])
            if not line["priced"]:
                for col in range(1, len(cols) + 1):
                    boq.cell(row=r, column=col).fill = WARN_FILL
            r += 1
        boq.cell(row=r, column=3, value=f"Total {trade}").font = BOLD
        c = boq.cell(row=r, column=7, value=f"=SUM(G{first}:G{r - 1})")
        c.number_format, c.font, c.border = MONEY, BOLD, TOP
        r += 2
    _widths(boq, [7, 22, 60, 7, 11, 12, 14, 11, 14, 40])
    boq.freeze_panes = "A6"

    # ---- Summary
    s = summary
    s["A1"] = "Tender Estimate Summary"
    s["A1"].font = Font(bold=True, size=14)
    info = [
        ("Project", project.get("name", "")),
        ("Client", project.get("client", "")),
        ("Site address", project.get("address", "")),
        ("Tender due", project.get("tender_date", "")),
        ("Location", f"{settings.location} (factor {settings.location_factor:g})"),
        ("Prepared", date.today().isoformat()),
    ]
    for i, (k, v) in enumerate(info, 3):
        s.cell(row=i, column=1, value=k).font = BOLD
        s.cell(row=i, column=2, value=str(v))

    row = 10
    _header(s, row, ["Trade", "", "Amount"])
    row += 1
    first_trade = row
    for trade in trades:
        s.cell(row=row, column=1, value=trade)
        s.cell(row=row, column=3, value=f"=SUMIF(BoQ!$B:$B,A{row},BoQ!$G:$G)").number_format = MONEY
        row += 1
    last_trade = row - 1
    row += 1

    def line(label, formula, pct=None, bold=False):
        nonlocal row
        s.cell(row=row, column=1, value=label).font = Font(bold=bold)
        if pct is not None:
            s.cell(row=row, column=2, value=pct).number_format = PCT
        c = s.cell(row=row, column=3, value=formula)
        c.number_format = MONEY
        if bold:
            c.font, c.border = BOLD, TOP
        row += 1
        return row - 1

    direct = line("Direct trade costs", f"=SUM(C{first_trade}:C{last_trade})" if trades else 0, bold=True)
    pre = line("Preliminaries", f"=C{direct}*B{row}/100", settings.preliminaries_pct)
    cont = line("Contingency", f"=(C{direct}+C{pre})*B{row}/100", settings.contingency_pct)
    esc = line("Escalation", f"=(C{direct}+C{pre})*B{row}/100", settings.escalation_pct)
    cost = line("Construction cost", f"=C{direct}+C{pre}+C{cont}+C{esc}", bold=True)
    oh = line("Overheads", f"=C{cost}*B{row}/100", settings.overhead_pct)
    pr = line("Profit", f"=C{cost}*B{row}/100", settings.profit_pct)
    ex = line("Tender price (excl. GST)", f"=C{cost}+C{oh}+C{pr}", bold=True)
    gst = line("GST", f"=C{ex}*B{row}/100", settings.gst_pct)
    line("Tender price (incl. GST)", f"=C{ex}+C{gst}", bold=True)
    row += 1
    unpriced = int((~lines["priced"]).sum()) if not lines.empty else 0
    notes = [
        "Percentages in column B can be edited; totals recalculate.",
        f"{unpriced} BoQ items have no rate (highlighted on the BoQ sheet)." if unpriced else "All BoQ items are priced.",
        "Prototype output - quantities and rates must be checked by an estimator before tender submission.",
    ]
    for n in notes:
        s.cell(row=row, column=1, value=n).font = Font(italic=True, color="555555")
        row += 1
    _widths(s, [40, 10, 18])

    # ---- Takeoff (raw)
    _header(raw, 1, list(takeoff.columns))
    for i, rec in enumerate(takeoff.itertuples(index=False), 2):
        for j, v in enumerate(rec, 1):
            raw.cell(row=i, column=j, value=v)
    _widths(raw, [10, 22, 60, 7, 11, 14, 40, 11, 11, 50])
    raw.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
